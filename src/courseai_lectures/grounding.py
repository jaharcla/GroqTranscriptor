"""Course-scoped grounding for ASR prompting and conservative transcript review."""

import json
import logging
import os
import time
from pathlib import Path

log = logging.getLogger(__name__)

TEXT_BLOCK_TYPES = {
    "paragraph",
    "heading_1",
    "heading_2",
    "heading_3",
    "bulleted_list_item",
    "numbered_list_item",
    "to_do",
    "toggle",
    "quote",
    "callout",
}
LOCAL_TYPES = {".md", ".txt", ".pdf", ".pptx"}


def _plain(rich_text):
    return "".join(
        item.get("plain_text", item.get("text", {}).get("content", ""))
        for item in rich_text or []
    ).strip()


def _page_title(page):
    for prop in page.get("properties", {}).values():
        if prop.get("type") == "title":
            title = _plain(prop.get("title", []))
            if title:
                return title
    return page.get("id", "Untitled Notion page")


def _block_text(block):
    kind = block.get("type")
    if kind == "child_page":
        return block.get("child_page", {}).get("title", "").strip()
    if kind not in TEXT_BLOCK_TYPES:
        return ""
    return _plain(block.get(kind, {}).get("rich_text", []))


def _material_text(path, limit):
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8", errors="replace")[:limit]
    if suffix == ".pdf":
        from pypdf import PdfReader

        parts = []
        total = 0
        for page in PdfReader(str(path)).pages:
            text = page.extract_text() or ""
            if text:
                parts.append(text)
                total += len(text)
            if total >= limit:
                break
        return "\n".join(parts)[:limit]
    if suffix == ".pptx":
        from pptx import Presentation

        parts = []
        total = 0
        for index, slide in enumerate(Presentation(str(path)).slides, start=1):
            slide_text = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    text = shape.text.strip()
                    if text:
                        slide_text.append(text)
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                        if cells:
                            slide_text.append(" | ".join(cells))
            if slide_text:
                value = f"[Slide {index}] " + "\n".join(slide_text)
                parts.append(value)
                total += len(value)
            if total >= limit:
                break
        return "\n".join(parts)[:limit]
    return ""


class CourseGrounder:
    """Build a reproducible context bundle scoped to one already-pinned course."""

    def __init__(self, config, notion):
        self.config = config
        self.notion = notion

    def _cache_path(self, course):
        return self.config.review_dir / "course-context" / f"{course}.json"

    def _add(self, sources, source_id, source_type, title, text, budget):
        text = " ".join(str(text or "").split())
        if not text or budget[0] <= 0:
            return
        text = text[: budget[0]]
        sources.append(
            {
                "id": str(source_id),
                "type": source_type,
                "title": str(title or source_id),
                "text": text,
            }
        )
        budget[0] -= len(text)

    def _walk_blocks(self, parent_id, label, sources, budget, visited, depth=0):
        if depth > 4 or budget[0] <= 0 or parent_id in visited:
            return
        visited.add(parent_id)
        for block in self.notion.children(parent_id):
            if budget[0] <= 0:
                return
            block_id = block.get("id")
            kind = block.get("type", "block")
            text = _block_text(block)
            if text:
                self._add(
                    sources,
                    f"notion:{block_id}",
                    f"notion_{kind}",
                    label,
                    text,
                    budget,
                )
            if block.get("has_children") and block_id:
                child_label = (
                    block.get("child_page", {}).get("title")
                    if kind == "child_page"
                    else label
                )
                self._walk_blocks(
                    block_id,
                    child_label or label,
                    sources,
                    budget,
                    visited,
                    depth + 1,
                )

    def _remote_context(self, course):
        root_id = self.config.courses.get(course)
        if not root_id:
            raise ValueError(f"Missing course mapping: {course}")
        max_chars = int(getattr(self.config, "notion_context_max_chars", 18000))
        budget = [max_chars]
        sources = []
        root = self.notion.request("GET", f"pages/{root_id}")
        root_title = _page_title(root)
        self._add(
            sources,
            f"notion-page:{root_id}",
            "notion_course_page",
            root_title,
            root_title,
            budget,
        )

        relation_names = {
            name.casefold()
            for name in getattr(
                self.config,
                "notion_context_relations",
                ("Modules", "Syllabus"),
            )
        }
        related = []
        for name, prop in root.get("properties", {}).items():
            if name.casefold() not in relation_names or prop.get("type") != "relation":
                continue
            related.extend(item.get("id") for item in prop.get("relation", []) if item.get("id"))

        visited = set()
        self._walk_blocks(root_id, root_title, sources, budget, visited)
        for page_id in dict.fromkeys(related):
            if budget[0] <= 0:
                break
            page = self.notion.request("GET", f"pages/{page_id}")
            title = _page_title(page)
            self._add(
                sources,
                f"notion-page:{page_id}",
                "notion_course_relation",
                title,
                title,
                budget,
            )
            self._walk_blocks(page_id, title, sources, budget, visited)
        return {
            "course": course,
            "course_page_id": root_id,
            "retrieved_at": time.time(),
            "sources": sources,
            "warnings": [],
        }

    def _cached_remote(self, course):
        cache = self._cache_path(course)
        ttl = float(getattr(self.config, "notion_context_cache_minutes", 30)) * 60
        if cache.exists():
            try:
                data = json.loads(cache.read_text(encoding="utf-8"))
                age = time.time() - float(data.get("retrieved_at", 0))
                if age <= ttl:
                    return data
            except (OSError, ValueError, TypeError):
                pass

        try:
            data = self._remote_context(course)
            cache.parent.mkdir(parents=True, exist_ok=True)
            temporary = cache.with_suffix(".tmp")
            temporary.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
            os.replace(temporary, cache)
            return data
        except Exception as exc:
            log.warning("Notion grounding unavailable for %s: %s", course, exc)
            if cache.exists():
                try:
                    data = json.loads(cache.read_text(encoding="utf-8"))
                    data.setdefault("warnings", []).append(
                        "Using stale cached Notion course context because refresh failed."
                    )
                    return data
                except (OSError, ValueError, TypeError):
                    pass
            return {
                "course": course,
                "course_page_id": self.config.courses.get(course, ""),
                "retrieved_at": None,
                "sources": [],
                "warnings": ["Notion course context was unavailable for this review."],
            }

    def _local_context(self, lecture, budget):
        sources = []
        manual = self.config.context_dir / f"{lecture.course}.md"
        if manual.exists():
            self._add(
                sources,
                f"local-context:{manual.name}",
                "local_course_context",
                manual.name,
                manual.read_text(encoding="utf-8", errors="replace"),
                budget,
            )

        materials_dir = getattr(self.config, "materials_dir", None)
        if not materials_dir:
            return sources
        course_dir = Path(materials_dir) / lecture.course
        if not course_dir.is_dir():
            return sources
        candidates = sorted(
            (
                path
                for path in course_dir.iterdir()
                if path.is_file() and path.suffix.lower() in LOCAL_TYPES
            ),
            key=lambda path: path.stat().st_mtime,
            reverse=True,
        )[:3]
        for path in candidates:
            if budget[0] <= 0:
                break
            try:
                text = _material_text(path, budget[0])
            except Exception as exc:
                log.warning("Cannot read lecture material %s: %s", path, exc)
                continue
            self._add(
                sources,
                f"material:{path.name}",
                "local_lecture_material",
                path.name,
                text,
                budget,
            )
        return sources

    def context(self, lecture):
        max_chars = int(getattr(self.config, "notion_context_max_chars", 18000))
        budget = [max_chars]
        sources = self._local_context(lecture, budget)
        warnings = []
        course_page_id = self.config.courses.get(lecture.course, "")
        if getattr(self.config, "notion_context_enabled", True) and budget[0] > 0:
            remote = self._cached_remote(lecture.course)
            course_page_id = remote.get("course_page_id", course_page_id)
            warnings.extend(remote.get("warnings", []))
            for item in remote.get("sources", []):
                if budget[0] <= 0:
                    break
                self._add(
                    sources,
                    item.get("id", "notion"),
                    item.get("type", "notion"),
                    item.get("title", "Notion"),
                    item.get("text", ""),
                    budget,
                )
        return {
            "course": lecture.course,
            "course_page_id": course_page_id,
            "sources": sources,
            "warnings": warnings,
        }

    def asr_prompt(self, lecture, context):
        """Return a <=224-token-ish spelling/context hint for Whisper."""
        titles = []
        snippets = []
        for source in context.get("sources", []):
            title = source.get("title", "").strip()
            if title and title not in titles:
                titles.append(title)
            text = source.get("text", "").strip()
            if text and text != title:
                snippets.append(text)
        prompt = (
            f"University lecture for {lecture.course}. "
            f"Course references: {'; '.join(titles[:8])}. "
            f"Use these spellings/terms when supported by the audio: {' '.join(snippets)[:650]}"
        )
        return prompt[:900]
